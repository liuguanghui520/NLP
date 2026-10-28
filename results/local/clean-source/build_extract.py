from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '方向C_项目定位评分标准与交付物.docx'
URL = 'https://6ccfkuh9.qwenwork.host/course-project/课程项目任务书_20261013_codex.html'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Microsoft YaHei'
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_after = Pt(6)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.13
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def h(text, level=2):
    doc.add_heading(text, level=level)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    for ridx, row in enumerate(t.rows):
        # Allow long cells to expand, but keep each requirement together across pages.
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if ridx == 0:
            props.append(OxmlElement('w:tblHeader'))
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DCE6F1' if ridx == 0 else ('F5F7FA' if ridx % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement('w:' + edge)
                for k, v in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    b.set(qn('w:' + k), v)
                borders.append(b)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + edge)
                m.set(qn('w:w'), '95')
                m.set(qn('w:type'), 'dxa')
                margins.append(m)
            tcpr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.08
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    run.bold = ridx == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

doc.add_heading('方向 C 课程项目要求摘录', 0)
p('自然语言处理课程  2026—2027 秋季学期')
p('已选方向：提示 vs 微调对比。项目须围绕完整的小型 NLP 流水线展开，包含非大模型基线、大模型方案以及两者的对比评估。以下按项目定位、评分标准、交付物及相关报告要求整理。')
h('一 项目定位', 1)
p('课程：自然语言处理；授课对象：软件工程专业三年级本科生；主讲：郭贵冰。项目共 45 分，占课程总评 45%；2—4 人一组，建议 3—4 人；不设答辩。')
h('通用定位')
p('用一条完整的小型 NLP 流水线，把课程里学过的方法真正用起来：数据 → 方法 → 评估 → 错误分析。外壳是大模型应用，内核是课程方法。每个项目必须包含：① 一个非大模型基线；② 一个大模型方案；③ 两者的对比评估。')
p('评分关注四件事：有没有基线对比、有没有评估数据、有没有把失败讲清楚、别人能不能复现。任务书强调：7B 模型加完整评估与诚实分析，优于使用 70B 模型却只给出一个分数。')
h('方向 C 的明确要求')
table(['条目','网页要求'], [
('方向名称','C. 提示 vs 微调对比'),
('大模型方案','零样本 / 少样本提示、LoRA 轻量微调'),
('参考基线','TF-IDF、n-gram 或小模型全参微调'),
('主要覆盖讲次','第 5、7、12、13 讲')], [3.2,13.8])
p('选题鼓励结合金融科技或智慧健康；也可自选题目。选题确认时须写清问题定义、成功判据，以及与课程方法的对应关系。')
h('课程原理的实验支撑')
p('每个项目的报告必须用自己实验的证据解释至少 3 个课程原理点；只复述概念不算，必须有实验或数据支撑。任务书可选原理点完整列于末页。')

doc.add_page_break()
h('二 评分标准', 1)
p('项目总分 45 分，评分表适用于方向 C。')
table(['评分项','分值','评分要点'], [
('选题与方案设计','5','问题定义清楚、成功判据可检验、技术路线有依据、分工到人。'),
('进度与工作量证明','5','有可核查的推进证据：数据就绪、基线跑通、迭代记录、运行日志 / 截图。'),
('源码与可复现性','12','一键可跑、README 完整、依赖与种子齐全、教师 / AI 复现通过、代码可读。'),
('实验与评估分析','12','有基线对比、指标口径清楚、≥3 类错误归因、≥1 组参数对比、结论有数据支撑。'),
('报告与演示录屏','6','报告结构规范、图表完整（4 分）；录屏真实可跑通、含难例（2 分）。'),
('个人贡献与工作量核验','5','每人一份贡献声明，模块、产出证据、工作量自评齐全且相互一致（3 分）；报告第⑧节分工与实际产出对得上（2 分）。')], [4.1,1.2,11.7])
h('两条硬门槛')
p('跑不通＝项目不及格：教师或 AI 按 README 无法复现主结果的，允许 48 小时内补交环境说明或修复；补交后仍无法跑通的，第③项记 0 分，总分不超过 25 分。')
p('报告与代码 / 日志对不上：报告中的关键数字在代码、日志或输出文件中找不到出处的，第④项减半；编造实验数据按学术不诚信处理。')
h('三个拉开分差的地方')
p('① 主动写出失败案例与局限，至少 3 个，逐一归因。\n② 有参数对比实验并给出结论，任务书举例包括 top-k、chunk 大小、LoRA 秩、不同模型。\n③ 能说清成本与延迟，包括 token 花费、响应时间。')
h('逾期扣分')
p('提交截止：2026-10-30（周五）24:00。以下扣分从项目总分 45 分中扣除。', True)
table(['网页列示的逾期时长','扣分'], [
('<24 小时（10-31 24:00 前）','−1'),
('<48 小时（11-01 24:00 前）','−2'),
('<72 小时（11-02 24:00 前）','−3'),
('≥72 小时','−5')], [13.5,3.5])

doc.add_page_break()
h('三 交付物与提交规范', 1)
table(['交付物','内容要求','提交形式'], [
('① 源码与 README','README 须含环境、依赖、一条可执行命令、预期输出、随机种子；同时提交完整代码。','同一个 zip 包'),
('② 项目报告','按八节结构撰写，正文 8—12 页，不含附录；格式为 PDF。','同一个 zip 包'),
('③ 演示录屏','≤2 分钟；真实跑一遍系统，至少展示一个难例；画面中可见本人学号与姓名。','填网盘链接'),
('④ 个人贡献声明','每人单独一份；写清负责模块、产出证据、工作量自评、对队友的确认；模板见附件。','同一个 zip 包'),
('⑤ AI 使用声明','工具名、使用环节、如何验证 AI 输出；可写在报告附录。','包含在报告内'),
('⑥ 项目摘要','300 字以内；包含问题、方法、主要结果（给出具体数字）、结论；不含图表。','同一个 zip 包；文件名为 项目摘要.md')], [3.2,10.5,3.3])
h('压缩包规范')
p('格式与数量：仅限 zip，只交 1 个压缩包。不收 rar、7z、分卷压缩，不要把材料拆成多个包。')
p('体积与外链：单个文件 ≤50 MB；模型权重、录屏、数据集一律不进包，只给链接。')
p('命名：组长学号_姓名_项目名.zip。网页示例：20246130_张三_公告要素抽取.zip。')
p('包内结构建议：报告.pdf、项目摘要.md、README.md、src/（全部源码）、贡献声明_学号_姓名.md（每人一份）。')
p('严禁打进包：虚拟环境、checkpoint、API Key、个人隐私数据。')
h('提交渠道与截止时间')
p('金山文档收集表：https://f.kdocs.cn/g/5htZQGSM')
p('截止：2026-10-30（周五）24:00。截止前可修改，以最后一次提交为准。')
h('任务书提供的附件')
p('项目报告模板_20261013_codex.md：八节结构的写作骨架，可直接复制填写。\n贡献声明表_20261013_codex.md：每人单独填写一份。\n两项附件位于原网页同目录。')

doc.add_page_break()
h('四 报告结构与图表要求', 1)
p('项目报告正文 8—12 页，不含附录；以下为任务书各节建议篇幅。')
table(['章节','内容要点','建议篇幅'], [
('① 问题定义与背景','解决什么问题、给谁用、成功判据是什么、为什么值得做。','0.5—1 页'),
('② 相关工作与方案选型','对比哪些替代方案、为什么选择该方案；须回应“为什么不用微调 / 为什么不用 RAG”这类问题。','1—1.5 页'),
('③ 系统设计','架构图（必配）、模块划分、数据流。','1—1.5 页'),
('④ 实现细节','关键技术点、踩过的坑与解法；这部分最能体现工作量。','2—3 页'),
('⑤ 评估','测试集来源与规模、指标定义、结果表、与基线对比、参数对比实验。','2—3 页'),
('⑥ 错误分析与局限','失败样例 ≥3 个，逐一归因，如数据噪声、过拟合、指令歧义、检索失败等。','1—2 页'),
('⑦ 工程与伦理','成本、延迟、权限与隐私、数据许可、AI 使用声明。','0.5—1 页'),
('⑧ 分工与贡献','每人负责的模块与产出证据，与个人贡献声明一致。','0.5 页')], [4.5,10,2.5])
p('图表要求：坐标轴、单位、图例齐全；loss / 指标曲线必须来自真实运行日志。')
h('数据与工具的相关要求')
p('数据优先使用公开数据，报告须写明来源、规模、许可与清洗规则。网页参考数据包括 THUCNews、ChnSentiCorp、MSRA、人民日报、CLUENER2020、CMRC2018、CBLUE 类基准任务和 SQuAD；使用前须自行确认许可与下载可用性。金融方向可从公开披露渠道构造数据。')
p('允许调用外部 API，费用自负；报告写明所用模型、调用量、单价与总花费估算。API Key 使用环境变量，不得写进代码或提交仓库，提交前自查。不得将隐私或未授权个人数据上传外部服务；医学、金融数据必须脱敏。')
p('允许本地小模型 / CPU 方案。预算或算力受限可缩小数据规模，在报告中说明降级原因与影响，不因此扣分。')
h('AI 使用与贡献声明')
p('鼓励用 AI 查文档、调试代码、生成样板、润色文字、做对比实验、互查错误。禁止用 AI 代写实验设计与结果分析，禁止编造数据、伪造日志或截图、抄袭代码与报告，以及将 API Key、隐私数据交给外部服务或写入提交物。')
p('报告必须声明 AI 的使用范围与人工修改部分。每人的贡献声明须写明负责模块与对应产出证据，与报告第⑧节一致；贡献声明相互矛盾且无法解释的，第⑥评分项相应扣分。')

doc.add_page_break()
h('五 可选课程原理点与来源', 1)
p('报告须从以下原理点中至少解释 3 个，每个均须结合本组真实实验或数据。')
table(['课程原理点','讲次'], [
('分词 / 子词与词表选择对检索或抽取效果的影响','第 2 讲'),
('词向量与句向量的表示能力差异','第 5 讲'),
('N-gram 与神经网络语言模型的取舍、困惑度（PPL）与生成质量的关系','第 7—8 讲'),
('序列标注的标签约束与实体级评测口径','第 9 讲'),
('注意力机制为什么能缓解信息瓶颈','第 10 讲'),
('Transformer 的位置编码、缩放点积与复杂度取舍','第 11 讲'),
('预训练目标（AR 与 MLM）对下游任务的影响','第 12 讲'),
('SFT / LoRA / 提示的本质差别与适用边界','第 13—14 讲'),
('RAG 的“召回是上限”、重排与引用校验的作用','第 15 讲'),
('评测指标与评测协议的选择依据','第 14—15 讲')], [13.5,3.5])
h('网页来源')
p('《自然语言处理》课程项目任务书；网页标示公布日期：2026-10-16（第 2 次随堂测验当天）。提取日期：2026-10-18。')
# A real hyperlink keeps the long source URL usable without printing percent-encoded text.
para = doc.add_paragraph('原始任务书：')
link = OxmlElement('w:hyperlink')
rid = para.part.relate_to(URL, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
link.set(qn('r:id'), rid)
r = OxmlElement('w:r')
t = OxmlElement('w:t')
t.text = '打开课程项目任务书网页'
r.append(t)
link.append(r)
para._p.append(link)

doc.core_properties.title = '方向 C 课程项目要求摘录'
doc.core_properties.subject = '项目定位、评分标准、交付物与报告要求'
doc.save(OUT)
print(OUT)
